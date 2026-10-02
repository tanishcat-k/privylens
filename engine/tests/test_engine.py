import io
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from fastapi.testclient import TestClient
from api.index import app
from api.privylens.engine import apply_text_redaction, find_pii

SAMPLE = """Email: tanishka@example.com
Phone: 9876543210
Card: 4111 1111 1111 1111
password = hunter2hunter
AWS: AKIAIOSFODNN7EXAMPLE
Link: https://app.site.io/cb?access_token=abcd1234efgh&x=1
Aadhaar: 2345 6789 0124 PAN: ABCPD1234E
Order id 1234 5678 9012 3456"""

def types(t): return [f["type"] for f in find_pii(t)]

def test_core_types():
    t = types(SAMPLE)
    for want in ["EMAIL", "PHONE", "CARD", "PASSWORD", "AWS_KEY", "SENSITIVE_URL", "PAN"]:
        assert want in t, (want, t)

def test_luhn_rejects_random_numbers():
    assert "CARD" not in types("Order id 1234 5678 9012 3456")

def test_placeholder_password_ignored():
    assert "PASSWORD" not in types("password: ********")

def test_redaction_text():
    out = apply_text_redaction("Email: a@b.com", find_pii("Email: a@b.com"))
    assert out == "Email: [REDACTED]"

def test_docx_roundtrip():
    d = Document(); p = d.add_paragraph(); p.add_run("Mail me at tan").bold = True; p.add_run("ishka@example.com now")
    b = io.BytesIO(); d.save(b)
    c = TestClient(app)
    r = c.post("/api/scan", files={"file": ("a.docx", b.getvalue())})
    assert r.status_code == 200 and r.json()["findings"][0]["type"] == "EMAIL"
    r = c.post("/api/redact", files={"file": ("a.docx", b.getvalue())})
    assert "example.com" not in "".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)

def test_image_roundtrip():
    img = Image.new("RGB", (900, 220), "white"); d = ImageDraw.Draw(img)
    f = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 30)
    d.text((20, 20), "Email: tanishka@example.com", fill="black", font=f)
    d.text((20, 90), "Phone: 9876543210", fill="black", font=f)
    b = io.BytesIO(); img.save(b, "PNG")
    c = TestClient(app)
    j = c.post("/api/scan", files={"file": ("a.png", b.getvalue())}).json()
    assert {"EMAIL", "PHONE"} <= {x["type"] for x in j["findings"]}, j["findings"]
    r = c.post("/api/redact", files={"file": ("a.png", b.getvalue())})
    out = Image.open(io.BytesIO(r.content)).convert("L")
    assert out.getextrema()[0] == 0 and sum(1 for p in out.getdata() if p == 0) > 3000
    r2 = c.post("/api/redact", files={"file": ("a.png", b.getvalue())}, data={"exclude": '["p0-0","p0-1"]'})
    assert r2.content != r.content
